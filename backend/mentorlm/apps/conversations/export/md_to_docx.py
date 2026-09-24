"""Markdown ответа ИИ → абзацы .docx с оформлением по ГОСТ (см. gost.py).

Разбор делает `markdown-it-py` (reference-порт CommonMark), а здесь — только
обход потока токенов. Свой парсер писать не стали сознательно: блочный автомат
плюс инлайн-токенизатор с правилом delimiter-run — это сотни строк и длинный
хвост краевых случаев, а цена ошибки высокая, потому что получившийся файл
человек несёт преподавателю.

Поток у markdown-it плоский: инлайн приходит последовательностью
`strong_open / text / strong_close`, что ложится на последовательность `run`-ов
со стеком форматов — рекурсия по инлайну не нужна.

Чего сознательно НЕ поддерживаем: сырой HTML (выводим как обычный текст — в
.docx он всё равно ничего не значит, а интерпретировать вывод модели как
разметку мы не станем нигде) и картинки (файлов у нас нет, остаётся подпись).
"""

from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .gost import (
    CODE_STYLE,
    FIRST_LINE_INDENT_CM,
    LIST_INDENT_STEP_CM,
    TABLE_FONT_SIZE_PT,
    apply_font,
    repeat_table_header,
    shade,
)

# Потолки на случай странного ввода: таблица на тысячу строк или список на
# двадцать уровней — это не документ, а способ уронить сборку.
MAX_TABLE_ROWS = 400
MAX_LIST_DEPTH = 6

# Схемы ссылок, которые разрешаем класть в документ. `relate_to` пишет строку в
# XML как есть, и `javascript:` из ответа модели стал бы проблемой уже той
# программы, которая откроет файл. Проверка та же, что во фронтовом safeUrl.
_SAFE_SCHEMES = ("http://", "https://", "mailto:")

_LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
_CODE_SHADE = "F2F2F2"


def _md():
    """Парсер: CommonMark плюс таблицы и зачёркивание.

    Не «gfm-like»: тот тянет linkify-it-py ради автолинковки, а модели ставят
    ссылки явно.
    """
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark").enable(["table", "strikethrough"])


def _safe_href(url: str) -> str | None:
    """Адрес, который можно записать в документ, либо None."""
    clean = (url or "").strip()
    low = clean.lower()
    return clean if any(low.startswith(s) for s in _SAFE_SCHEMES) else None


def _horizontal_rule(paragraph) -> None:
    """Линия-разделитель: нижняя граница пустого абзаца (в API её нет)."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


class _Renderer:
    """Обход плоского потока токенов с накоплением состояния списков и цитат."""

    def __init__(self, document):
        self.doc = document
        # Стек открытых списков: (маркированный ли, счётчик номера).
        self.lists: list[list] = []
        self.quote_depth = 0
        # Маркер, который должен уйти в начало следующего абзаца списка.
        self.pending_marker: str | None = None

    # ── Абзацы ────────────────────────────────────────────────────────────

    def _paragraph(self, style: str | None = None):
        """Новый абзац с учётом текущей вложенности списка и цитаты."""
        par = self.doc.add_paragraph(style=style)
        fmt = par.paragraph_format
        depth = len(self.lists) + self.quote_depth
        if depth:
            fmt.left_indent = Cm(LIST_INDENT_STEP_CM * depth)
            # Внутри списка красная строка мешает: маркер должен начинать строку.
            fmt.first_line_indent = Cm(0)
        if self.quote_depth:
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return par

    def _take_marker(self, par) -> None:
        """Выдать накопленный маркер списка первым run'ом абзаца."""
        if self.pending_marker is None:
            return
        run = par.add_run(self.pending_marker)
        self.pending_marker = None

    # ── Инлайн ────────────────────────────────────────────────────────────

    def _add_run(self, par, text: str, fmt: dict):
        run = par.add_run(text)
        run.bold = fmt.get("bold", False)
        run.italic = fmt.get("italic", False)
        run.font.strike = fmt.get("strike", False)
        if fmt.get("code"):
            from .gost import CODE_FONT_NAME, CODE_FONT_SIZE_PT

            run.font.name = CODE_FONT_NAME
            run.font.size = Pt(CODE_FONT_SIZE_PT)
            apply_font(run._r, CODE_FONT_NAME)
        return run

    def _add_link(self, par, url: str, text: str) -> None:
        """Внешняя ссылка: в python-docx её нет, нужен r:id из отношений части."""
        href = _safe_href(url)
        if href is None:
            # Небезопасную схему показываем текстом — сведений не теряем, а
            # кликабельной её не делаем.
            self._add_run(par, text, {})
            return
        rel_id = par.part.relate_to(href, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        run = par.add_run(text)
        run.font.color.rgb = _LINK_COLOR
        run.font.underline = True
        link.append(run._r)
        par._p.append(link)

    def _inline(self, par, token) -> None:
        """Разложить инлайн-поток в набор run'ов."""
        fmt = {"bold": False, "italic": False, "strike": False, "code": False}
        children = token.children or []
        i = 0
        while i < len(children):
            child = children[i]
            kind = child.type

            if kind == "text":
                self._add_run(par, child.content, fmt)
            elif kind == "code_inline":
                self._add_run(par, child.content, {**fmt, "code": True})
            elif kind == "strong_open":
                fmt["bold"] = True
            elif kind == "strong_close":
                fmt["bold"] = False
            elif kind == "em_open":
                fmt["italic"] = True
            elif kind == "em_close":
                fmt["italic"] = False
            elif kind == "s_open":
                fmt["strike"] = True
            elif kind == "s_close":
                fmt["strike"] = False
            elif kind == "link_open":
                # Текст ссылки собираем целиком до link_close: форматирование
                # внутри ссылки встречается редко, а так гиперссылка остаётся
                # одним run'ом и не разваливается на части.
                depth, j, parts = 1, i + 1, []
                while j < len(children) and depth:
                    if children[j].type == "link_open":
                        depth += 1
                    elif children[j].type == "link_close":
                        depth -= 1
                        if not depth:
                            break
                    elif children[j].type in ("text", "code_inline"):
                        parts.append(children[j].content)
                    j += 1
                url = child.attrGet("href") or ""
                self._add_link(par, url, "".join(parts) or url)
                i = j
            elif kind == "image":
                alt = child.content or child.attrGet("alt") or ""
                self._add_run(par, f"[изображение: {alt}]", {**fmt, "italic": True})
            elif kind == "softbreak":
                self._add_run(par, " ", fmt)
            elif kind == "hardbreak":
                par.add_run().add_break()
            elif kind in ("html_inline",):
                # Сырой HTML в .docx смысла не имеет — показываем как текст.
                self._add_run(par, child.content, fmt)
            i += 1

    # ── Блоки ─────────────────────────────────────────────────────────────

    def _fence(self, token) -> None:
        """Блок кода: каждая строка — отдельный абзац стиля MLM Code.

        Мягкие переносы внутри одного абзаца выглядели бы так же, но ломают и
        заливку, и перенос блока на следующую страницу.
        """
        body = (token.content or "").rstrip("\n")
        for line in body.split("\n"):
            par = self.doc.add_paragraph(style=CODE_STYLE)
            # Пробелы в начале строки Word схлопывать не станет, но пустая
            # строка без run'а теряет заливку — поэтому кладём хотя бы пустой.
            par.add_run(line or " ")
            shade(par, _CODE_SHADE)

    def _table(self, tokens, start: int) -> int:
        """Собрать таблицу; вернуть индекс токена после table_close."""
        rows: list[list[tuple[str, object]]] = []
        aligns: list[str | None] = []
        i = start + 1
        row: list[tuple[str, object]] | None = None

        while i < len(tokens) and tokens[i].type != "table_close":
            token = tokens[i]
            if token.type == "tr_open":
                row = []
            elif token.type == "tr_close":
                if row is not None and len(rows) < MAX_TABLE_ROWS:
                    rows.append(row)
                row = None
            elif token.type in ("th_open", "td_open"):
                style = token.attrGet("style") or ""
                if token.type == "th_open":
                    aligns.append(
                        "right"
                        if "right" in style
                        else "center" if "center" in style else None
                    )
                # Следом всегда идёт inline (или сразу закрытие у пустой ячейки).
                content = tokens[i + 1] if tokens[i + 1].type == "inline" else None
                if row is not None:
                    row.append((token.type, content))
            i += 1

        if rows:
            self._render_table(rows, aligns)
        return i + 1

    def _render_table(self, rows, aligns) -> None:
        cols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=0, cols=cols)
        table.style = "Table Grid"  # видимые границы — требование стандарта
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for index, row in enumerate(rows):
            cells = table.add_row().cells
            for col in range(cols):
                cell = cells[col]
                cell.text = ""
                par = cell.paragraphs[0]
                fmt = par.paragraph_format
                fmt.first_line_indent = Cm(0)
                fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                if col < len(aligns) and aligns[col] == "right":
                    fmt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif col < len(aligns) and aligns[col] == "center":
                    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER

                kind, content = row[col] if col < len(row) else ("td_open", None)
                if content is not None:
                    self._inline(par, content)
                for run in par.runs:
                    run.font.size = Pt(TABLE_FONT_SIZE_PT)
                    if kind == "th_open":
                        run.bold = True

            if index == 0:
                repeat_table_header(table.rows[0])

    # ── Главный цикл ──────────────────────────────────────────────────────

    def run(self, tokens) -> None:
        i = 0
        while i < len(tokens):
            token = tokens[i]
            kind = token.type

            if kind == "heading_open":
                level = int(token.tag[1])
                par = self.doc.add_heading("", level=min(level, 4))
                if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                    self._inline(par, tokens[i + 1])
                    i += 1
            elif kind == "paragraph_open":
                par = self._paragraph()
                self._take_marker(par)
                if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                    self._inline(par, tokens[i + 1])
                    i += 1
                if self.quote_depth:
                    for run in par.runs:
                        run.italic = True
            elif kind == "fence" or kind == "code_block":
                self._fence(token)
            elif kind == "bullet_list_open":
                if len(self.lists) < MAX_LIST_DEPTH:
                    self.lists.append([True, 0])
            elif kind == "ordered_list_open":
                if len(self.lists) < MAX_LIST_DEPTH:
                    self.lists.append([False, 0])
            elif kind in ("bullet_list_close", "ordered_list_close"):
                if self.lists:
                    self.lists.pop()
            elif kind == "list_item_open":
                # Маркеры ставим руками: нумерация Word живёт в numbering.xml,
                # которого в шаблоне python-docx для нужных уровней нет, а ГОСТ
                # 7.32 сам предписывает перечисления через дефис.
                if self.lists:
                    bullet, counter = self.lists[-1]
                    if bullet:
                        self.pending_marker = "— "
                    else:
                        start = int(token.info) if token.info.isdigit() else counter + 1
                        self.lists[-1][1] = start
                        self.pending_marker = f"{start}) "
            elif kind == "blockquote_open":
                self.quote_depth += 1
            elif kind == "blockquote_close":
                self.quote_depth = max(0, self.quote_depth - 1)
            elif kind == "table_open":
                i = self._table(tokens, i)
                continue
            elif kind == "hr":
                _horizontal_rule(self.doc.add_paragraph())
            elif kind == "html_block":
                par = self._paragraph()
                par.add_run(token.content.strip())

            i += 1


def render_markdown(document, text: str) -> None:
    """Дописать Markdown-текст в документ абзацами."""
    _Renderer(document).run(_md().parse(text or ""))
