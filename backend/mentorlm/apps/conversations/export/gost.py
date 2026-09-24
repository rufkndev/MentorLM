"""Оформление документа по ГОСТ 7.32-2017: стили, поля, колонтитул.

Единственное место, где живут числа стандарта. Требования конкретного вуза
обычно отличаются на одну-две величины (чаще всего кегль), поэтому они собраны
константами наверху — правка должна быть однострочной, а не поиском по коду.

Титульный лист не генерируем: его состав вуз задаёт сам, и подставить туда
что-то наугад хуже, чем не подставлять ничего.

Про правку XML. Всё, что делается через python-docx, сделано через python-docx.
Прямой XML остаётся там, где у библиотеки просто нет API: поле номера страницы,
четыре слота гарнитуры в w:rFonts, заливка абзаца и повтор шапки таблицы.
"""

from __future__ import annotations

from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── Числа стандарта ───────────────────────────────────────────────────────────
FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 14  # ГОСТ требует «не менее 12»; 14 — обычное требование вузов
CODE_FONT_NAME = "Consolas"
CODE_FONT_SIZE_PT = 11
TABLE_FONT_SIZE_PT = 12  # в таблице стандарт допускает кегль меньше основного

MARGIN_LEFT_CM = 3.0
MARGIN_RIGHT_CM = 1.5
MARGIN_TOP_CM = 2.0
MARGIN_BOTTOM_CM = 2.0

FIRST_LINE_INDENT_CM = 1.25
LIST_INDENT_STEP_CM = 1.25

# Стиль для блоков кода — своё имя, чтобы не конфликтовать со встроенными.
CODE_STYLE = "MLM Code"


def apply_font(element, name: str) -> None:
    """Прописать гарнитуру во все четыре слота w:rFonts.

    `Font.name` в python-docx заполняет только w:ascii и w:hAnsi. Для кириллицы
    этого обычно достаточно (Word относит её к hAnsi), но w:cs и w:eastAsia
    остаются пустыми, и в LibreOffice или «Р7-Офис» на смешанном тексте
    гарнитура подменяется. Это единственная причина лезть сюда в XML.
    """
    rpr = element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(slot), name)


def _add_page_number(paragraph) -> None:
    """Вставить поле { PAGE } — у python-docx полей нет вовсе.

    Три fldChar (begin/separate/end) плюс instrText — это то, как Word хранит
    вычисляемое поле. Текст между separate и end — кэш значения на случай, если
    открывающая программа поля не пересчитывает.
    """
    run = paragraph.add_run()._r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    run.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.append(separate)

    cached = OxmlElement("w:t")
    cached.text = "1"
    run.append(cached)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)


def shade(paragraph, fill: str) -> None:
    """Залить абзац цветом (подложка блока кода) — тоже только через XML."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def repeat_table_header(row) -> None:
    """Повторять строку-шапку на каждой странице (w:tblHeader)."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _setup_normal(document) -> None:
    """Базовый стиль текста: гарнитура, кегль, интервал, отступ, выравнивание."""
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    apply_font(normal.element, FONT_NAME)

    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = Cm(FIRST_LINE_INDENT_CM)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # ГОСТ не знает «интервала после абзаца»: расстояние задаёт межстрочный.
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def _setup_headings(document) -> None:
    """Заголовки: встроенные стили Word синие и на Calibri — переопределяем."""
    for level in range(1, 5):
        style = document.styles[f"Heading {level}"]
        style.font.name = FONT_NAME
        style.font.size = Pt(FONT_SIZE_PT)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        apply_font(style.element, FONT_NAME)

        fmt = style.paragraph_format
        # Заголовок первого уровня считаем структурным элементом — по центру и
        # без абзацного отступа; разделы ниже идут с отступа, как основной текст.
        fmt.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )
        fmt.first_line_indent = Cm(0 if level == 1 else FIRST_LINE_INDENT_CM)
        fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
        # Заголовок не должен остаться внизу страницы без своего текста.
        fmt.keep_with_next = True
        style.next_paragraph_style = document.styles["Normal"]


def _setup_code_style(document) -> None:
    """Стиль блока кода: моноширинный, одинарный интервал, без отступа строки."""
    style = document.styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = CODE_FONT_NAME
    style.font.size = Pt(CODE_FONT_SIZE_PT)
    style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    apply_font(style.element, CODE_FONT_NAME)

    fmt = style.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.first_line_indent = Cm(0)
    fmt.left_indent = Cm(FIRST_LINE_INDENT_CM)
    # Код по ширине не растягиваем: выравнивание раздуло бы пробелы в отступах.
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def build_document():
    """Пустой документ с оформлением по ГОСТ: поля, стили, номера страниц."""
    from docx import Document

    document = Document()

    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.left_margin = Cm(MARGIN_LEFT_CM)
    section.right_margin = Cm(MARGIN_RIGHT_CM)
    section.top_margin = Cm(MARGIN_TOP_CM)
    section.bottom_margin = Cm(MARGIN_BOTTOM_CM)

    _setup_normal(document)
    _setup_headings(document)
    _setup_code_style(document)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Без обнуления отступа «центр» уехал бы вправо на величину абзацного.
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    _add_page_number(footer)

    return document
