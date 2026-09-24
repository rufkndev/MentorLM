"""Экспорт ответа ИИ в файл: .docx с оформлением по ГОСТ и .md как есть.

`gost.py` — числа и стили стандарта, `md_to_docx.py` — обход Markdown.
Наружу торчат две сборки: ответ и список литературы.
"""

from __future__ import annotations

from io import BytesIO

from urllib.parse import urlsplit

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from .gost import build_document
from .md_to_docx import render_markdown

__all__ = ["build_docx", "build_sources_docx"]


def build_docx(blocks: list[tuple[str, str]]) -> BytesIO:
    """Собрать .docx из пар «заголовок, markdown» и вернуть готовый буфер.

    Список, а не одна строка, — чтобы выгрузка целого диалога («Вопрос» /
    «Ответ» вперемешку) не потребовала переписывать сигнатуру. Пустой заголовок
    просто не выводится.
    """
    document = build_document()
    for index, (title, body) in enumerate(blocks):
        if index:
            document.add_paragraph()
        if title:
            document.add_heading(title, level=1)
        render_markdown(document, body)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _domain(url: str) -> str:
    """Домен без www — он же название ресурса в библиографической записи."""
    host = urlsplit(url).netloc.lower()
    return host[4:] if host.startswith("www.") else host


def _gost_entry(source: dict, accessed: str) -> str:
    """Одна запись списка литературы по ГОСТ Р 7.0.5-2008 (электронный ресурс).

    Схема для ресурса удалённого доступа: заглавие, сведения о ресурсе, адрес и
    дата обращения. Заглавие берём из цитаты; если его нет — остаётся домен,
    потому что запись без заглавия стандарту не соответствует вовсе.
    """
    site = _domain(source.get("url", ""))
    title = (source.get("title") or "").strip() or site
    tail = f" // {site} : сайт." if site and title != site else ""
    return f"{title}{tail} — URL: {source.get('url', '')} (дата обращения: {accessed})."


def build_sources_docx(
    sources: list[dict], accessed: str, heading: str = "Список литературы"
):
    """Список литературы по ГОСТ: нумерованный, с датой обращения.

    Дата приходит из сообщения (день, когда модель читала страницы), а не
    берётся текущей: в записи по стандарту именно дата обращения, и подставлять
    сюда день выгрузки было бы неправдой.
    """
    document = build_document()
    document.add_heading(heading, level=1)

    for index, source in enumerate(sources, 1):
        par = document.add_paragraph()
        fmt = par.paragraph_format
        # Запись висячая: номер начинает строку, перенос идёт с отступом.
        fmt.first_line_indent = Cm(-0.75)
        fmt.left_indent = Cm(0.75 + 1.25)
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.add_run(f"{index}. {_gost_entry(source, accessed)}")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
